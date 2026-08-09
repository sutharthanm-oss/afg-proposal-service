"""
Generic proposal-generation engine.
Takes a template_id (which selects templates/<id>/master.pptx + coords.json)
and a data payload (matching the extraction-prompt.md schema), and produces
a filled PPTX. Adding a new product/org template means adding a new folder
under templates/ with its own master.pptx + coords.json -- no code changes.
"""
import json
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), "templates")


def _color(hex_str):
    hex_str = hex_str.lstrip("#")
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


def _px2in(px, scale):
    return Inches(px * scale)


def _add_fixed_textbox(slide, x_in, y_in, w_in, h_in, text, size_pt, color_hex,
                        bold=True, align=PP_ALIGN.LEFT, font_family="Calibri",
                        anchor=MSO_ANCHOR.MIDDLE, word_wrap=True):
    """Textbox with wrap='square' + noAutofit -- renders identically across
    LibreOffice, PowerPoint, and PDF viewers (see: alignment bug fixed earlier
    caused by python-pptx's default wrap='none' + spAutoFit)."""
    tb = slide.shapes.add_textbox(Inches(x_in), Inches(y_in), Inches(w_in), Inches(h_in))
    tf = tb.text_frame
    tf.word_wrap = word_wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = _color(color_hex)
    run.font.name = font_family
    return tb


def generate_proposal(template_id: str, data: dict, output_path: str) -> str:
    """
    data must match the extraction-prompt.md schema:
    {
      "prospect": {"name", "dob", "age", "smoking_status"},
      "tiers": [ {tier_label, monthly_premium, death, tpd, ... , remarks}, ... ],
      "agent_name": str   # pulled from Airtable by the caller, not from Claude extraction
    }
    """
    tpl_dir = os.path.join(TEMPLATES_DIR, template_id)
    with open(os.path.join(tpl_dir, "coords.json")) as f:
        cfg = json.load(f)

    prs = Presentation(os.path.join(tpl_dir, cfg["master_file"]))
    scale = cfg["canvas"]["scale_in_per_px"]
    font_family = cfg["font"]["family"]

    prospect = data["prospect"]
    agent_name = data.get("agent_name", "-")
    tiers = {t["tier_label"]: t for t in data["tiers"]}

    # ---------- TITLE SLIDE ----------
    ts = cfg["title_slide"]
    slide1 = prs.slides[ts["index"]]

    nb = ts["prospect_name_box"]
    _add_fixed_textbox(slide1, nb["x_in"], nb["y_in"], nb["w_in"], nb["h_in"],
                        prospect["name"], nb["size_pt"], nb["color"],
                        font_family=font_family)

    pb = ts["prepared_by_box"]
    _add_fixed_textbox(slide1, pb["x_in"], pb["y_in"], pb["w_in"], pb["h_in"],
                        f"Prepared By: {agent_name}", pb["size_pt"], pb["color"],
                        font_family=font_family)

    # ---------- QUOTE SLIDE ----------
    qs = cfg["quote_slide"]
    slide2 = prs.slides[qs["index"]]
    ib = qs["info_boxes"]

    def info_box(key, value):
        b = ib[key]
        x = b["x0"] * scale
        y = b["y0"] * scale
        w = (b["x1"] - b["x0"]) * scale
        h = (b["y1"] - b["y0"]) * scale
        _add_fixed_textbox(slide2, x, y, w, h, value, b["size_pt"], "002642",
                            align=PP_ALIGN.CENTER, font_family=font_family)

    info_box("age", prospect.get("age", "-"))
    info_box("class", "-")  # never available per extraction-prompt.md
    info_box("birthday", prospect.get("dob", "-"))
    info_box("smoking", prospect.get("smoking_status", "-"))

    cols = qs["columns"]

    def col_box(col_letter, y0, y1, value, size_pt, color_hex="1A1A1A"):
        x0, x1 = cols[col_letter]
        x = x0 * scale
        y = y0 * scale
        w = (x1 - x0) * scale
        h = (y1 - y0) * scale
        _add_fixed_textbox(slide2, x, y, w, h, value, size_pt, color_hex,
                            align=PP_ALIGN.CENTER, font_family=font_family)

    mp = qs["monthly_premium_row"]
    for col in ("A", "B", "C"):
        tier = tiers.get(col)
        val = tier["monthly_premium"] if tier else "-"
        col_box(col, mp["y0"], mp["y1"], val, mp["size_pt"])

    row_size = qs["data_row_size_pt"]
    for row in qs["data_rows"]:
        for col in ("A", "B", "C"):
            tier = tiers.get(col)
            val = tier.get(row["field"], "-") if tier else "-"
            col_box(col, row["y0"], row["y1"], val, row_size)

    rm = qs["remarks_row"]
    for col in ("A", "B", "C"):
        tier = tiers.get(col)
        val = tier["remarks"] if tier else "-"
        col_box(col, rm["y0"], rm["y1"], val, rm["size_pt"], color_hex=rm["color"])

    prs.save(output_path)
    return output_path


def list_templates():
    """Returns available template_ids -- used by the API to validate requests
    and by the bot to know which products it can currently generate for."""
    if not os.path.isdir(TEMPLATES_DIR):
        return []
    return [d for d in os.listdir(TEMPLATES_DIR)
            if os.path.isfile(os.path.join(TEMPLATES_DIR, d, "coords.json"))]
